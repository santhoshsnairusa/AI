import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()

# Define NexAware Colors
BG_COLOR = RGBColor(15, 23, 42)      # Dark Slate (Background)
TEXT_COLOR = RGBColor(248, 250, 252) # Off-white (Body Text)
ACCENT_COLOR = RGBColor(99, 102, 241) # Indigo (Titles/Accents)
SECONDARY_ACCENT = RGBColor(56, 189, 248) # Light Blue

def apply_nexaware_style(slide):
    # Set dark background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    # Style shapes
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                # If it's a title placeholder, use accent color
                if shape == slide.shapes.title:
                    run.font.color.rgb = ACCENT_COLOR
                    run.font.bold = True
                    run.font.name = 'Segoe UI'
                else:
                    run.font.color.rgb = TEXT_COLOR
                    run.font.name = 'Segoe UI'

def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    apply_nexaware_style(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    subtitle.text = subtitle_text
    
    # Re-apply styles since text assignment overwrites runs
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT_COLOR
    title.text_frame.paragraphs[0].runs[0].font.name = 'Segoe UI'
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    
    subtitle.text_frame.paragraphs[0].runs[0].font.color.rgb = SECONDARY_ACCENT
    subtitle.text_frame.paragraphs[0].runs[0].font.name = 'Segoe UI'
    
    return slide

def add_content_slide(prs, title_text, bullet_points):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    apply_nexaware_style(slide)
    
    title = slide.shapes.title
    body_shape = slide.placeholders[1]
    
    title.text = title_text
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT_COLOR
    title.text_frame.paragraphs[0].runs[0].font.name = 'Segoe UI'
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    
    tf = body_shape.text_frame
    tf.text = bullet_points[0]
    tf.paragraphs[0].runs[0].font.color.rgb = TEXT_COLOR
    tf.paragraphs[0].runs[0].font.name = 'Segoe UI'
    
    for point in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = point
        p.runs[0].font.color.rgb = TEXT_COLOR
        p.runs[0].font.name = 'Segoe UI'
        
    return slide

# Build the Slides
add_title_slide(prs, "NexAware Home AI", "Your Private, Intelligent Home Companion")

add_content_slide(prs, "The Problem", [
    "Modern homes are complex, and managing them is a headache:",
    "• Lost Information: Appliance manuals, paint codes, and WiFi passwords are scattered or lost.",
    "• Missed Deadlines: Warranties expire and filters go unchanged because we forget to track them.",
    "• Privacy Risks: Existing AI assistants send your most intimate home data to the cloud."
])

add_content_slide(prs, "The Solution: NexAware", [
    "NexAware is a 100% offline, local-first AI assistant that memorizes your home's unique data.",
    "• Private: Your data never leaves your network.",
    "• Intelligent: It reads and understands your specific manuals and notes.",
    "• Proactive: It tracks your home infrastructure and warns you before things expire."
])

add_content_slide(prs, "Feature 1: Smart Knowledge Base", [
    "Say goodbye to the junk drawer.",
    "• Upload Anything: Drag and drop PDF service manuals, text notes, or warranty receipts.",
    "• Advanced RAG: NexAware instantly indexes every word into a highly optimized vector database.",
    "• Instant Answers: Just ask, 'How do I change the filter on my LG fridge?' and get the exact steps instantly."
])

add_content_slide(prs, "Feature 2: Asset & Warranty Tracking", [
    "Protect your investments automatically.",
    "• Digital Inventory: Track the make, model, and brand of every appliance and system in your house.",
    "• Smart Deadlines: Log warranty end dates and product expiry dates.",
    "• Proactive Alerts: The NexAware dashboard automatically flags items 30 days before they expire."
])

add_content_slide(prs, "The Technology", [
    "Built on an enterprise-grade, modern technology stack:",
    "• Local LLM: Powered by lightweight, hyper-fast local models (like Llama 3) running via Ollama.",
    "• Vector Search: Utilizes pgvector for instant, semantic document retrieval.",
    "• Sleek Interface: A responsive, glassmorphism-inspired React dashboard.",
    "• Containerized: Deploys instantly via Docker Compose."
])

add_content_slide(prs, "Why NexAware?", [
    "NexAware Home AI vs Cloud Assistants (Alexa/Google):",
    "• Privacy: 100% Local & Private vs Data sold & analyzed",
    "• Knowledge: Personalized to your home vs Generic web answers",
    "• Cost: Free & Open vs Hidden subscriptions",
    "• Offline: Works offline always vs Fails without internet"
])

add_title_slide(prs, "Join the Future of Home Management", "Take back control of your home's data today.\nSmart. Local. Yours.")

# Save presentation
prs.save('NexAware_Pitch_Deck.pptx')
print("Successfully generated NexAware_Pitch_Deck.pptx")
